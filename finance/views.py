# finance/views.py

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import authenticate
from django.contrib.auth.models import User
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404

from rest_framework import generics, status
from rest_framework.permissions import IsAuthenticated, AllowAny
from rest_framework.response import Response
from rest_framework.views import APIView

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from openpyxl import Workbook

from .models import (
    Profile,
    Income,
    Expense,
    Budget,
    SavingsGoal,
    Notification,
    Report,
)

from .serializers import (
    ProfileSerializer,
    IncomeSerializer,
    ExpenseSerializer,
    BudgetSerializer,
    SavingsGoalSerializer,
    NotificationSerializer,
    ReportSerializer,
    RegisterSerializer,
)


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def parse_date(value):
    if not value:
        return None

    return datetime.strptime(
        value,
        "%Y-%m-%d"
    ).date()


def get_month_dates(month_string):
    year, month = map(
        int,
        month_string.split("-")
    )

    start_date = date(
        year,
        month,
        1
    )

    if month == 12:
        end_date = date(
            year + 1,
            1,
            1
        )
    else:
        end_date = date(
            year,
            month + 1,
            1
        )

    return start_date, end_date


# ============================================================
# REGISTER
# ============================================================

class RegisterView(generics.CreateAPIView):
    serializer_class = RegisterSerializer
    permission_classes = [AllowAny]


# ============================================================
# PROFILE
# ============================================================

class ProfileView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        profile, created = Profile.objects.get_or_create(
            user=request.user
        )

        serializer = ProfileSerializer(profile)

        return Response(
            serializer.data,
            status=status.HTTP_200_OK
        )

    def put(self, request):
        profile, created = Profile.objects.get_or_create(
            user=request.user
        )

        serializer = ProfileSerializer(
            profile,
            data=request.data,
            partial=True
        )

        if serializer.is_valid():
            serializer.save(
                user=request.user
            )

            return Response(
                serializer.data,
                status=status.HTTP_200_OK
            )

        return Response(
            serializer.errors,
            status=status.HTTP_400_BAD_REQUEST
        )


class MyProfileView(ProfileView):
    pass


# ============================================================
# CHANGE PASSWORD
# ============================================================

class ChangePasswordView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):

        old_password = request.data.get(
            "old_password"
        )

        new_password = request.data.get(
            "new_password"
        )

        if not old_password or not new_password:
            return Response(
                {
                    "success": False,
                    "message": (
                        "old_password and "
                        "new_password are required."
                    ),
                },
                status=status.HTTP_400_BAD_REQUEST
            )

        user = authenticate(
            username=request.user.username,
            password=old_password
        )

        if user is None:
            return Response(
                {
                    "success": False,
                    "message": "Old password is incorrect.",
                },
                status=status.HTTP_400_BAD_REQUEST
            )

        if len(new_password) < 8:
            return Response(
                {
                    "success": False,
                    "message": (
                        "New password must be at least "
                        "8 characters long."
                    ),
                },
                status=status.HTTP_400_BAD_REQUEST
            )

        request.user.set_password(
            new_password
        )

        request.user.save()

        return Response(
            {
                "success": True,
                "message": "Password changed successfully.",
            },
            status=status.HTTP_200_OK
        )


# ============================================================
# BUDGET ALERT HELPERS
# ============================================================

# ============================================================
# BUDGET ALERT HELPERS
# ============================================================

def create_budget_notification(
    user,
    budget,
    percentage
):
    """
    Create or update a budget notification
    when utilization reaches or exceeds 80%.
    """

    if percentage >= 100:

        title = "Budget Exceeded"

        message = (
            f"Your {budget.category} budget has "
            f"been exceeded. "
            f"Current utilization is "
            f"{percentage:.1f}%."
        )

    else:

        title = "Budget Limit Alert"

        message = (
            f"Your {budget.category} budget is "
            f"{percentage:.1f}% utilized. "
            f"You have reached the 80% "
            f"warning threshold."
        )

    existing = Notification.objects.filter(
        user=user,
        notification_type="budget",
        related_id=budget.id,
        title=title,
        is_read=False,
    ).first()

    if existing:

        existing.message = message

        existing.save(
            update_fields=["message"]
        )

        return

    Notification.objects.create(
        user=user,
        notification_type="budget",
        title=title,
        message=message,
        related_id=budget.id,
    )


def check_budget_alerts(
    user,
    expense
):
    """
    Check active budgets matching
    the expense category.
    """

    budgets = Budget.objects.filter(
        user=user,
        category__iexact=expense.category,
        start_date__lte=expense.date,
        end_date__gte=expense.date,
    )

    for budget in budgets:

        total_expense = Expense.objects.filter(
            user=user,
            category__iexact=budget.category,
            date__gte=budget.start_date,
            date__lte=budget.end_date,
        ).aggregate(
            total=Sum("amount")
        )["total"] or Decimal("0")

        if budget.amount > 0:

            percentage = (
                total_expense /
                budget.amount
            ) * Decimal("100")

            if percentage >= 80:

                create_budget_notification(
                    user,
                    budget,
                    float(percentage)
                )

# ============================================================
# EXPENSES
# ============================================================

class MyExpensesView(generics.ListAPIView):
    serializer_class = ExpenseSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Expense.objects.filter(
            user=self.request.user
        ).order_by("-date", "-id")


class ExpenseCreateView(
    generics.ListCreateAPIView
):
    serializer_class = ExpenseSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Expense.objects.filter(
            user=self.request.user
        ).order_by("-date", "-id")

    def perform_create(self, serializer):

        expense = serializer.save(
            user=self.request.user
        )

        check_budget_alerts(
            self.request.user,
            expense
        )


class ExpenseDetailView(
    generics.RetrieveUpdateDestroyAPIView
):
    serializer_class = ExpenseSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Expense.objects.filter(
            user=self.request.user
        )

    def perform_update(self, serializer):

        expense = serializer.save(
            user=self.request.user
        )

        check_budget_alerts(
            self.request.user,
            expense
        )


class ExpenseListCreateView(
    ExpenseCreateView
):
    pass


# ============================================================
# INCOME
# ============================================================

class MyIncomeView(generics.ListAPIView):
    serializer_class = IncomeSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Income.objects.filter(
            user=self.request.user
        ).order_by("-date", "-id")


class IncomeCreateView(
    generics.ListCreateAPIView
):
    serializer_class = IncomeSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Income.objects.filter(
            user=self.request.user
        ).order_by("-date", "-id")

    def perform_create(self, serializer):
        serializer.save(
            user=self.request.user
        )


class IncomeDetailView(
    generics.RetrieveUpdateDestroyAPIView
):
    serializer_class = IncomeSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Income.objects.filter(
            user=self.request.user
        )


class IncomeListCreateView(
    IncomeCreateView
):
    pass


# ============================================================
# BUDGETS
# ============================================================

class MyBudgetsView(generics.ListAPIView):
    serializer_class = BudgetSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Budget.objects.filter(
            user=self.request.user
        ).order_by("-id")


class BudgetCreateView(
    generics.ListCreateAPIView
):
    serializer_class = BudgetSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Budget.objects.filter(
            user=self.request.user
        ).order_by("-id")

    def perform_create(self, serializer):
        serializer.save(
            user=self.request.user
        )


class BudgetDetailView(
    generics.RetrieveUpdateDestroyAPIView
):
    serializer_class = BudgetSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Budget.objects.filter(
            user=self.request.user
        )


class BudgetListCreateView(
    BudgetCreateView
):
    pass


# ============================================================
# SAVINGS MILESTONE HELPER
# ============================================================

def check_savings_milestone(
    savings_goal
):
    """
    Create separate notifications when a savings
    goal reaches 80% or 100%.

    Prevents duplicate unread notifications
    for the same milestone.
    """

    if savings_goal.target_amount <= 0:
        return

    percentage = (
        savings_goal.current_amount /
        savings_goal.target_amount
    ) * Decimal("100")

    milestone = None

    if percentage >= 100:
        milestone = 100

    elif percentage >= 80:
        milestone = 80

    if milestone is None:
        return

    # --------------------------------------------------------
    # Decide notification
    # --------------------------------------------------------

    if milestone == 100:

        title = "Savings Goal Completed"

        message = (
            f"Congratulations! You completed "
            f"your '{savings_goal.name}' "
            f"savings goal."
        )

    else:

        title = "Savings Goal Milestone"

        message = (
            f"You have reached "
            f"{percentage:.1f}% of your "
            f"'{savings_goal.name}' "
            f"savings goal."
        )

    # --------------------------------------------------------
    # Duplicate check
    # --------------------------------------------------------

    existing = Notification.objects.filter(
        user=savings_goal.user,
        notification_type="savings",
        related_id=savings_goal.id,
        title=title,
        is_read=False,
    ).exists()

    if existing:
        return

    # --------------------------------------------------------
    # Create notification
    # --------------------------------------------------------

    Notification.objects.create(
        user=savings_goal.user,
        notification_type="savings",
        title=title,
        message=message,
        related_id=savings_goal.id,
    )


# ============================================================
# SAVINGS GOALS
# ============================================================

class MySavingsGoalsView(
    generics.ListAPIView
):
    serializer_class = SavingsGoalSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return SavingsGoal.objects.filter(
            user=self.request.user
        ).order_by("-id")


class SavingsGoalCreateView(
    generics.ListCreateAPIView
):
    serializer_class = SavingsGoalSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return SavingsGoal.objects.filter(
            user=self.request.user
        ).order_by("-id")

    def perform_create(self, serializer):

        savings_goal = serializer.save(
            user=self.request.user
        )

        if (
            savings_goal.current_amount
            >= savings_goal.target_amount
        ):

            savings_goal.status = "completed"

        else:

            savings_goal.status = "active"

        savings_goal.save(
            update_fields=["status"]
        )

        check_savings_milestone(
            savings_goal
        )


class SavingsGoalDetailView(
    generics.RetrieveUpdateDestroyAPIView
):
    serializer_class = SavingsGoalSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return SavingsGoal.objects.filter(
            user=self.request.user
        )

    def perform_update(self, serializer):

        savings_goal = serializer.save(
            user=self.request.user
        )

        if (
            savings_goal.current_amount
            >= savings_goal.target_amount
        ):

            savings_goal.status = "completed"

        else:

            savings_goal.status = "active"

        savings_goal.save(
            update_fields=["status"]
        )

        check_savings_milestone(
            savings_goal
        )


class SavingsGoalListCreateView(
    SavingsGoalCreateView
):
    pass


# ============================================================
# NOTIFICATIONS
# ============================================================

class MyNotificationsView(
    generics.ListAPIView
):
    serializer_class = NotificationSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Notification.objects.filter(
            user=self.request.user
        ).order_by("-created_at")


class MyUnreadNotificationsView(
    generics.ListAPIView
):
    serializer_class = NotificationSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Notification.objects.filter(
            user=self.request.user,
            is_read=False,
        ).order_by("-created_at")


class NotificationMarkReadView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):

        notification = get_object_or_404(
            Notification,
            id=pk,
            user=request.user,
        )

        notification.is_read = True

        notification.save(
            update_fields=["is_read"]
        )

        return Response(
            {
                "success": True,
                "message": (
                    "Notification marked as read."
                ),
                "notification":
                    NotificationSerializer(
                        notification
                    ).data,
            },
            status=status.HTTP_200_OK,
        )


class NotificationReadView(
    NotificationMarkReadView
):
    pass


class MarkAllNotificationsReadView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):

        updated = Notification.objects.filter(
            user=request.user,
            is_read=False,
        ).update(
            is_read=True
        )

        return Response(
            {
                "success": True,
                "message": (
                    "All notifications marked as read."
                ),
                "updated_count": updated,
            },
            status=status.HTTP_200_OK,
        )


# ============================================================
# ANALYTICS
# ============================================================

class AnalyticsView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):

        user = request.user

        month = request.query_params.get(
            "month"
        )

        start_date_param = request.query_params.get(
            "start_date"
        )

        end_date_param = request.query_params.get(
            "end_date"
        )

        start_date = None
        end_date = None

        # ----------------------------------------------------
        # MONTH FILTER
        # ----------------------------------------------------

        if month:

            try:

                start_date, end_date = (
                    get_month_dates(month)
                )

            except (
                ValueError,
                TypeError
            ):

                return Response(
                    {
                        "success": False,
                        "message": (
                            "Invalid month format. "
                            "Use YYYY-MM."
                        ),
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

        # ----------------------------------------------------
        # DATE RANGE FILTER
        # ----------------------------------------------------

        elif (
            start_date_param
            or end_date_param
        ):

            try:

                start_date = parse_date(
                    start_date_param
                )

                end_date = parse_date(
                    end_date_param
                )

            except ValueError:

                return Response(
                    {
                        "success": False,
                        "message": (
                            "Invalid date format. "
                            "Use YYYY-MM-DD."
                        ),
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

            if (
                start_date
                and end_date
                and start_date > end_date
            ):

                return Response(
                    {
                        "success": False,
                        "message": (
                            "start_date cannot be "
                            "after end_date."
                        ),
                    },
                    status=status.HTTP_400_BAD_REQUEST,
                )

        # ----------------------------------------------------
        # USER DATA
        # ----------------------------------------------------

        incomes = Income.objects.filter(
            user=user
        )

        expenses = Expense.objects.filter(
            user=user
        )

        if start_date:

            incomes = incomes.filter(
                date__gte=start_date
            )

            expenses = expenses.filter(
                date__gte=start_date
            )

        if end_date:

            if month:

                last_day = date.fromordinal(
                    end_date.toordinal() - 1
                )

                incomes = incomes.filter(
                    date__lte=last_day
                )

                expenses = expenses.filter(
                    date__lte=last_day
                )

            else:

                incomes = incomes.filter(
                    date__lte=end_date
                )

                expenses = expenses.filter(
                    date__lte=end_date
                )

        # ----------------------------------------------------
        # SUMMARY
        # ----------------------------------------------------

        total_income = (
            incomes.aggregate(
                total=Sum("amount")
            )["total"]
            or Decimal("0")
        )

        total_expenses = (
            expenses.aggregate(
                total=Sum("amount")
            )["total"]
            or Decimal("0")
        )

        balance = (
            total_income -
            total_expenses
        )

        # ----------------------------------------------------
        # CATEGORY-WISE EXPENSES
        # ----------------------------------------------------

        categories = [
            "Food",
            "Travel",
            "Shopping",
            "Education",
            "Entertainment",
            "Miscellaneous",
        ]

        category_data = []

        for category in categories:

            category_total = (
                expenses.filter(
                    category__iexact=category
                ).aggregate(
                    total=Sum("amount")
                )["total"]
                or Decimal("0")
            )

            category_data.append(
                {
                    "category": category,
                    "amount": float(
                        category_total
                    ),
                }
            )

        # ----------------------------------------------------
        # MONTHLY TRENDS
        # ----------------------------------------------------

        monthly_trends = []

        if month:

            trend_start = start_date
            trend_end = end_date

            current_year = (
                trend_start.year
            )

            current_month = (
                trend_start.month
            )

            while True:

                current_start = date(
                    current_year,
                    current_month,
                    1
                )

                if current_month == 12:

                    current_end = date(
                        current_year + 1,
                        1,
                        1
                    )

                else:

                    current_end = date(
                        current_year,
                        current_month + 1,
                        1
                    )

                if current_start >= trend_end:
                    break

                month_income = (
                    Income.objects.filter(
                        user=user,
                        date__gte=current_start,
                        date__lt=current_end,
                    ).aggregate(
                        total=Sum("amount")
                    )["total"]
                    or Decimal("0")
                )

                month_expense = (
                    Expense.objects.filter(
                        user=user,
                        date__gte=current_start,
                        date__lt=current_end,
                    ).aggregate(
                        total=Sum("amount")
                    )["total"]
                    or Decimal("0")
                )

                monthly_trends.append(
                    {
                        "month":
                            current_start.strftime(
                                "%Y-%m"
                            ),
                        "income":
                            float(month_income),
                        "expense":
                            float(month_expense),
                        "balance":
                            float(
                                month_income -
                                month_expense
                            ),
                    }
                )

                if current_month == 12:

                    current_year += 1
                    current_month = 1

                else:

                    current_month += 1

        else:

            months = {}

            for income in Income.objects.filter(
                user=user
            ):

                key = income.date.strftime(
                    "%Y-%m"
                )

                if key not in months:

                    months[key] = {
                        "income": Decimal("0"),
                        "expense": Decimal("0"),
                    }

                months[key]["income"] += (
                    income.amount
                )

            for expense in Expense.objects.filter(
                user=user
            ):

                key = expense.date.strftime(
                    "%Y-%m"
                )

                if key not in months:

                    months[key] = {
                        "income": Decimal("0"),
                        "expense": Decimal("0"),
                    }

                months[key]["expense"] += (
                    expense.amount
                )

            for key in sorted(months):

                month_income = (
                    months[key]["income"]
                )

                month_expense = (
                    months[key]["expense"]
                )

                monthly_trends.append(
                    {
                        "month": key,
                        "income":
                            float(month_income),
                        "expense":
                            float(month_expense),
                        "balance":
                            float(
                                month_income -
                                month_expense
                            ),
                    }
                )

        # ----------------------------------------------------
        # SAVINGS SUMMARY
        # ----------------------------------------------------

        savings_goals = SavingsGoal.objects.filter(
            user=user
        )

        total_target = (
            savings_goals.aggregate(
                total=Sum("target_amount")
            )["total"]
            or Decimal("0")
        )

        total_saved = (
            savings_goals.aggregate(
                total=Sum("current_amount")
            )["total"]
            or Decimal("0")
        )

        remaining = (
            total_target -
            total_saved
        )

        if remaining < 0:
            remaining = Decimal("0")

        savings_percentage = Decimal("0")

        if total_target > 0:

            savings_percentage = (
                total_saved /
                total_target
            ) * Decimal("100")

        return Response(
            {
                "success": True,

                "filters": {
                    "month": month,
                    "start_date":
                        start_date_param,
                    "end_date":
                        end_date_param,
                },

                "summary": {
                    "total_income":
                        float(total_income),
                    "total_expenses":
                        float(total_expenses),
                    "balance":
                        float(balance),
                },

                "category_wise_expenses":
                    category_data,

                "monthly_trends":
                    monthly_trends,

                "savings_summary": {
                    "total_target":
                        float(total_target),
                    "total_saved":
                        float(total_saved),
                    "remaining":
                        float(remaining),
                    "completion_percentage":
                        float(
                            savings_percentage
                        ),
                },
            },
            status=status.HTTP_200_OK,
        )


# ============================================================
# REPORTS LIST
# ============================================================

class MyReportsView(
    generics.ListAPIView
):
    serializer_class = ReportSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Report.objects.filter(
            user=self.request.user
        ).order_by("-created_at")


# ============================================================
# REPORT CREATE
# ============================================================

class ReportCreateView(
    generics.ListCreateAPIView
):
    serializer_class = ReportSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return Report.objects.filter(
            user=self.request.user
        ).order_by("-created_at")

    def perform_create(self, serializer):

        serializer.save(
            user=self.request.user
        )


# ============================================================
# MONTHLY REPORT GENERATION
# ============================================================

class GenerateMonthlyReportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):

        month = request.query_params.get(
            "month"
        )

        if not month:

            return Response(
                {
                    "success": False,
                    "message": (
                        "month parameter is required. "
                        "Use YYYY-MM."
                    ),
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:

            start_date, end_date = (
                get_month_dates(month)
            )

        except (
            ValueError,
            TypeError
        ):

            return Response(
                {
                    "success": False,
                    "message": (
                        "Invalid month format. "
                        "Use YYYY-MM."
                    ),
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        user = request.user

        incomes = Income.objects.filter(
            user=user,
            date__gte=start_date,
            date__lt=end_date
        )

        expenses = Expense.objects.filter(
            user=user,
            date__gte=start_date,
            date__lt=end_date
        )

        total_income = (
            incomes.aggregate(
                total=Sum("amount")
            )["total"]
            or Decimal("0")
        )

        total_expenses = (
            expenses.aggregate(
                total=Sum("amount")
            )["total"]
            or Decimal("0")
        )

        balance = (
            total_income -
            total_expenses
        )

        # ----------------------------------------------------
        # CATEGORY EXPENSES
        # ----------------------------------------------------

        categories = [
            "Food",
            "Travel",
            "Shopping",
            "Education",
            "Entertainment",
            "Miscellaneous",
        ]

        category_expenses = {}

        for category in categories:

            amount = (
                expenses.filter(
                    category__iexact=category
                ).aggregate(
                    total=Sum("amount")
                )["total"]
                or Decimal("0")
            )

            category_expenses[
                category
            ] = float(amount)

        # ----------------------------------------------------
        # BUDGETS
        # ----------------------------------------------------

        budgets = Budget.objects.filter(
            user=user,
            start_date__lt=end_date,
            end_date__gte=start_date
        )

        budget_data = []

        for budget in budgets:

            spent = (
                expenses.filter(
                    category__iexact=
                        budget.category
                ).aggregate(
                    total=Sum("amount")
                )["total"]
                or Decimal("0")
            )

            utilization = Decimal("0")

            if budget.amount > 0:

                utilization = (
                    spent /
                    budget.amount
                ) * Decimal("100")

            budget_data.append(
                {
                    "category":
                        budget.category,
                    "budget":
                        float(
                            budget.amount
                        ),
                    "spent":
                        float(spent),
                    "remaining":
                        float(
                            budget.amount -
                            spent
                        ),
                    "utilization_percentage":
                        float(
                            utilization
                        ),
                }
            )

        # ----------------------------------------------------
        # SAVINGS GOALS
        # ----------------------------------------------------

        savings_goals = SavingsGoal.objects.filter(
            user=user
        )

        savings_data = []

        for goal in savings_goals:

            remaining = (
                goal.target_amount -
                goal.current_amount
            )

            if remaining < 0:
                remaining = Decimal("0")

            percentage = Decimal("0")

            if goal.target_amount > 0:

                percentage = (
                    goal.current_amount /
                    goal.target_amount
                ) * Decimal("100")

            savings_data.append(
                {
                    "name":
                        goal.name,

                    "target_amount":
                        float(
                            goal.target_amount
                        ),

                    "current_amount":
                        float(
                            goal.current_amount
                        ),

                    "remaining":
                        float(remaining),

                    "completion_percentage":
                        float(percentage),

                    "status":
                        goal.status,
                }
            )

        # ----------------------------------------------------
        # CREATE REPORT RECORD
        # ----------------------------------------------------

        report = Report.objects.create(
            user=user,
            title=(
                f"{month} Financial Report"
            ),
            report_type="monthly",
            start_date=start_date,
            end_date=end_date,
        )

        return Response(
            {
                "success": True,

                "message": (
                    "Monthly report generated "
                    "successfully."
                ),

                "report": {
                    "id":
                        report.id,

                    "title":
                        report.title,

                    "report_type":
                        report.report_type,

                    "start_date":
                        start_date,

                    "end_date":
                        end_date,

                    "created_at":
                        report.created_at,
                },

                "summary": {
                    "total_income":
                        float(total_income),

                    "total_expenses":
                        float(total_expenses),

                    "balance":
                        float(balance),
                },

                "category_wise_expenses":
                    category_expenses,

                "budgets":
                    budget_data,

                "savings_goals":
                    savings_data,
            },
            status=status.HTTP_200_OK,
        )


# ============================================================
# PDF REPORT
# ============================================================

class GeneratePDFReportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):

        month = request.query_params.get(
            "month"
        )

        if not month:

            return Response(
                {
                    "success": False,
                    "message": (
                        "month parameter is required. "
                        "Use YYYY-MM."
                    ),
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:

            start_date, end_date = (
                get_month_dates(month)
            )

        except (
            ValueError,
            TypeError
        ):

            return Response(
                {
                    "success": False,
                    "message": (
                        "Invalid month format. "
                        "Use YYYY-MM."
                    ),
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        user = request.user

        incomes = Income.objects.filter(
            user=user,
            date__gte=start_date,
            date__lt=end_date
        )

        expenses = Expense.objects.filter(
            user=user,
            date__gte=start_date,
            date__lt=end_date
        )

        total_income = (
            incomes.aggregate(
                total=Sum("amount")
            )["total"]
            or Decimal("0")
        )

        total_expenses = (
            expenses.aggregate(
                total=Sum("amount")
            )["total"]
            or Decimal("0")
        )

        balance = (
            total_income -
            total_expenses
        )

        buffer = BytesIO()

        document = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40,
        )

        styles = getSampleStyleSheet()

        story = []

        story.append(
            Paragraph(
                f"BudgetBuddy - {month} Financial Report",
                styles["Title"]
            )
        )

        story.append(
            Spacer(1, 20)
        )

        story.append(
            Paragraph(
                f"User: {user.username}",
                styles["Normal"]
            )
        )

        story.append(
            Paragraph(
                f"Period: {start_date} to "
                f"{date.fromordinal(end_date.toordinal() - 1)}",
                styles["Normal"]
            )
        )

        story.append(
            Spacer(1, 20)
        )

        # ----------------------------------------------------
        # SUMMARY
        # ----------------------------------------------------

        summary_data = [
            ["Summary", "Amount"],
            [
                "Total Income",
                f"Rs.{total_income:.2f}"
            ],
            [
                "Total Expenses",
                f"Rs.{total_expenses:.2f}"
            ],
            [
                "Balance",
                f"Rs.{balance:.2f}"
            ],
        ]

        summary_table = Table(
            summary_data,
            colWidths=[
                3 * inch,
                2 * inch
            ]
        )

        summary_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.grey
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        1,
                        colors.black
                    ),
                    (
                        "ALIGN",
                        (1, 1),
                        (-1, -1),
                        "RIGHT"
                    ),
                    (
                        "PADDING",
                        (0, 0),
                        (-1, -1),
                        6
                    ),
                ]
            )
        )

        story.append(
            summary_table
        )

        story.append(
            Spacer(1, 25)
        )

        # ----------------------------------------------------
        # CATEGORY EXPENSES
        # ----------------------------------------------------

        story.append(
            Paragraph(
                "Category-wise Expenses",
                styles["Heading2"]
            )
        )

        category_data = [
            ["Category", "Amount"]
        ]

        categories = [
            "Food",
            "Travel",
            "Shopping",
            "Education",
            "Entertainment",
            "Miscellaneous",
        ]

        for category in categories:

            amount = (
                expenses.filter(
                    category__iexact=category
                ).aggregate(
                    total=Sum("amount")
                )["total"]
                or Decimal("0")
            )

            category_data.append(
                [
                    category,
                    f"Rs.{amount:.2f}"
                ]
            )

        category_table = Table(
            category_data,
            colWidths=[
                3 * inch,
                2 * inch
            ]
        )

        category_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.grey
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        1,
                        colors.black
                    ),
                    (
                        "ALIGN",
                        (1, 1),
                        (-1, -1),
                        "RIGHT"
                    ),
                    (
                        "PADDING",
                        (0, 0),
                        (-1, -1),
                        6
                    ),
                ]
            )
        )

        story.append(
            category_table
        )

        story.append(
            Spacer(1, 25)
        )

        # ----------------------------------------------------
        # BUDGET SUMMARY
        # ----------------------------------------------------

        story.append(
            Paragraph(
                "Budget Summary",
                styles["Heading2"]
            )
        )

        budgets = Budget.objects.filter(
            user=user,
            start_date__lt=end_date,
            end_date__gte=start_date
        )

        budget_data = [
            [
                "Category",
                "Budget",
                "Spent",
                "Utilization"
            ]
        ]

        for budget in budgets:

            spent = (
                expenses.filter(
                    category__iexact=
                        budget.category
                ).aggregate(
                    total=Sum("amount")
                )["total"]
                or Decimal("0")
            )

            utilization = Decimal("0")

            if budget.amount > 0:

                utilization = (
                    spent /
                    budget.amount
                ) * Decimal("100")

            budget_data.append(
                [
                    budget.category,
                    f"Rs.{budget.amount:.2f}",
                    f"Rs.{spent:.2f}",
                    f"{utilization:.1f}%"
                ]
            )

        if len(budget_data) == 1:

            budget_data.append(
                [
                    "No budgets",
                    "-",
                    "-",
                    "-"
                ]
            )

        budget_table = Table(
            budget_data,
            colWidths=[
                1.5 * inch,
                1.4 * inch,
                1.4 * inch,
                1.2 * inch
            ]
        )

        budget_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.grey
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        1,
                        colors.black
                    ),
                    (
                        "PADDING",
                        (0, 0),
                        (-1, -1),
                        5
                    ),
                ]
            )
        )

        story.append(
            budget_table
        )

        story.append(
            Spacer(1, 25)
        )

        # ----------------------------------------------------
        # SAVINGS GOALS
        # ----------------------------------------------------

        story.append(
            Paragraph(
                "Savings Goals",
                styles["Heading2"]
            )
        )

        savings_goals = SavingsGoal.objects.filter(
            user=user
        )

        savings_data = [
            [
                "Goal",
                "Target",
                "Saved",
                "Remaining"
            ]
        ]

        for goal in savings_goals:

            remaining = (
                goal.target_amount -
                goal.current_amount
            )

            if remaining < 0:
                remaining = Decimal("0")

            savings_data.append(
                [
                    goal.name,
                    f"Rs.{goal.target_amount:.2f}",
                    f"Rs.{goal.current_amount:.2f}",
                    f"Rs.{remaining:.2f}"
                ]
            )

        if len(savings_data) == 1:

            savings_data.append(
                [
                    "No savings goals",
                    "-",
                    "-",
                    "-"
                ]
            )

        savings_table = Table(
            savings_data,
            colWidths=[
                2 * inch,
                1.5 * inch,
                1.5 * inch,
                1.5 * inch
            ]
        )

        savings_table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.grey
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        1,
                        colors.black
                    ),
                    (
                        "PADDING",
                        (0, 0),
                        (-1, -1),
                        5
                    ),
                ]
            )
        )

        story.append(
            savings_table
        )

        story.append(
            Spacer(1, 20)
        )

        story.append(
            Paragraph(
                "Generated by BudgetBuddy",
                styles["Normal"]
            )
        )

        document.build(
            story
        )

        buffer.seek(0)

        response = HttpResponse(
            buffer,
            content_type="application/pdf"
        )

        response[
            "Content-Disposition"
        ] = (
            f'attachment; '
            f'filename="BudgetBuddy_{month}.pdf"'
        )

        return response


# ============================================================
# EXCEL REPORT
# ============================================================

class GenerateExcelReportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):

        month = request.query_params.get(
            "month"
        )

        if not month:

            return Response(
                {
                    "success": False,
                    "message": (
                        "month parameter is required. "
                        "Use YYYY-MM."
                    ),
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:

            start_date, end_date = (
                get_month_dates(month)
            )

        except (
            ValueError,
            TypeError
        ):

            return Response(
                {
                    "success": False,
                    "message": (
                        "Invalid month format. "
                        "Use YYYY-MM."
                    ),
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        user = request.user

        incomes = Income.objects.filter(
            user=user,
            date__gte=start_date,
            date__lt=end_date
        )

        expenses = Expense.objects.filter(
            user=user,
            date__gte=start_date,
            date__lt=end_date
        )

        total_income = (
            incomes.aggregate(
                total=Sum("amount")
            )["total"]
            or Decimal("0")
        )

        total_expenses = (
            expenses.aggregate(
                total=Sum("amount")
            )["total"]
            or Decimal("0")
        )

        balance = (
            total_income -
            total_expenses
        )

        workbook = Workbook()

        # ----------------------------------------------------
        # SUMMARY SHEET
        # ----------------------------------------------------

        summary_sheet = workbook.active

        summary_sheet.title = "Summary"

        summary_sheet.append(
            [
                "BudgetBuddy Monthly Financial Report"
            ]
        )

        summary_sheet.append(
            [
                "User",
                user.username
            ]
        )

        summary_sheet.append(
            [
                "Month",
                month
            ]
        )

        summary_sheet.append([])

        summary_sheet.append(
            [
                "Summary",
                "Amount"
            ]
        )

        summary_sheet.append(
            [
                "Total Income",
                float(total_income)
            ]
        )

        summary_sheet.append(
            [
                "Total Expenses",
                float(total_expenses)
            ]
        )

        summary_sheet.append(
            [
                "Balance",
                float(balance)
            ]
        )

        # ----------------------------------------------------
        # EXPENSE SHEET
        # ----------------------------------------------------

        expense_sheet = workbook.create_sheet(
            "Expenses"
        )

        expense_sheet.append(
            [
                "ID",
                "Category",
                "Amount",
                "Date",
                "Description"
            ]
        )

        for expense in expenses:

            expense_sheet.append(
                [
                    expense.id,
                    expense.category,
                    float(expense.amount),
                    expense.date,
                    expense.description,
                ]
            )

        # ----------------------------------------------------
        # INCOME SHEET
        # ----------------------------------------------------

        income_sheet = workbook.create_sheet(
            "Income"
        )

        income_sheet.append(
            [
                "ID",
                "Source",
                "Amount",
                "Date",
                "Description"
            ]
        )

        for income in incomes:

            income_sheet.append(
                [
                    income.id,
                    income.source,
                    float(income.amount),
                    income.date,
                    income.description,
                ]
            )

        # ----------------------------------------------------
        # BUDGET SHEET
        # ----------------------------------------------------

        budget_sheet = workbook.create_sheet(
            "Budgets"
        )

        budget_sheet.append(
            [
                "Category",
                "Budget",
                "Spent",
                "Remaining",
                "Utilization %"
            ]
        )

        budgets = Budget.objects.filter(
            user=user,
            start_date__lt=end_date,
            end_date__gte=start_date
        )

        for budget in budgets:

            spent = (
                expenses.filter(
                    category__iexact=
                        budget.category
                ).aggregate(
                    total=Sum("amount")
                )["total"]
                or Decimal("0")
            )

            utilization = Decimal("0")

            if budget.amount > 0:

                utilization = (
                    spent /
                    budget.amount
                ) * Decimal("100")

            budget_sheet.append(
                [
                    budget.category,
                    float(
                        budget.amount
                    ),
                    float(spent),
                    float(
                        budget.amount -
                        spent
                    ),
                    float(
                        utilization
                    ),
                ]
            )

        # ----------------------------------------------------
        # SAVINGS SHEET
        # ----------------------------------------------------

        savings_sheet = workbook.create_sheet(
            "Savings Goals"
        )

        savings_sheet.append(
            [
                "Goal",
                "Target Amount",
                "Current Amount",
                "Remaining",
                "Completion %",
                "Status"
            ]
        )

        savings_goals = SavingsGoal.objects.filter(
            user=user
        )

        for goal in savings_goals:

            remaining = (
                goal.target_amount -
                goal.current_amount
            )

            if remaining < 0:
                remaining = Decimal("0")

            percentage = Decimal("0")

            if goal.target_amount > 0:

                percentage = (
                    goal.current_amount /
                    goal.target_amount
                ) * Decimal("100")

            savings_sheet.append(
                [
                    goal.name,
                    float(
                        goal.target_amount
                    ),
                    float(
                        goal.current_amount
                    ),
                    float(remaining),
                    float(percentage),
                    goal.status,
                ]
            )

        # ----------------------------------------------------
        # AUTO COLUMN WIDTH
        # ----------------------------------------------------

        for sheet in workbook.worksheets:

            for column in sheet.columns:

                max_length = 0

                column_letter = (
                    column[0].column_letter
                )

                for cell in column:

                    if cell.value is not None:

                        length = len(
                            str(cell.value)
                        )

                        if length > max_length:
                            max_length = length

                sheet.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 2,
                    40
                )

        output = BytesIO()

        workbook.save(
            output
        )

        output.seek(0)

        response = HttpResponse(
            output,
            content_type=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            )
        )

        response[
            "Content-Disposition"
        ] = (
            f'attachment; '
            f'filename="BudgetBuddy_{month}.xlsx"'
        )

        return response


# ============================================================
# OPTIONAL REPORT STATUS
# ============================================================

class ReportGenerationStatusView(
    APIView
):
    permission_classes = [IsAuthenticated]

    def get(self, request):

        report_count = Report.objects.filter(
            user=request.user
        ).count()

        return Response(
            {
                "success": True,
                "reports_generated":
                    report_count,
                "status": "ready",
            },
            status=status.HTTP_200_OK,
        )


# ============================================================
# OPTIONAL FINANCE HEALTH CHECK
# ============================================================

class FinanceHealthCheckView(
    APIView
):
    permission_classes = [IsAuthenticated]

    def get(self, request):

        return Response(
            {
                "success": True,
                "message":
                    "Finance API is working.",
                "user":
                    request.user.username,
            },
            status=status.HTTP_200_OK,
        )